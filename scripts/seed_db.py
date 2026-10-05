"""
Builds the Firebase Realtime Database seed for VE-ARCHEO from the original
VE16ARCHEO spreadsheets. Replaces the 2016 Apps Script importer, which wrote
each row's data under the next row's key and dropped the last row.

Inputs:
  data/VE16ARCHEO-Site Data.xlsx   -> sites (SiteForm) and US layers (USLayers)
  data/VE16-ARCHEO-qGIS data.xlsx  -> QGIS map points

Output:
  data/seed.json  -> {"data": ..., "groups": ..., "gis": ...}

Keys are deterministic, so re-running and re-uploading replaces records
instead of duplicating them. Upload with a PATCH on the root so /editors is
kept, e.g.:
  firebase database:update / data/seed.json --project ve-archeo

Usage:
  pip install openpyxl
  python scripts/seed_db.py
"""
import json
import os
import re
import sys
from collections import Counter
from datetime import datetime

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SITE_XLSX = os.path.join(ROOT, "data", "VE16ARCHEO-Site Data.xlsx")
GIS_XLSX = os.path.join(ROOT, "data", "VE16-ARCHEO-qGIS data.xlsx")
OUT = os.path.join(ROOT, "data", "seed.json")

IMPORT_DATE = "2016-12-01"  # date of the original VE16ARCHEO import
RECORDER = "VE16ARCHEO (import)"

PLACEHOLDERS = {"", "none", "none given", "not given", "n/a"}

# Align the sheet's English/typo values with the Italian vocabulary used in QGIS.
EPOCHS = {
    "roman": "Eta Romana",
    "high medieval": "Altomedioevo",  # the team's "highmed" layer is Altomedioevo
    "iron age": "Eta Del Ferro",
    "late antiquity": "Tarda Antichita",
    "pre-romana": "Eta Preromana",
}
TYPES = {
    "repertp": "Reperto",
    "carotaggi": "Carotaggi",
    "brick": "Muro In Mattoni",
    "beick": "Muro In Mattoni",
    "tree": "Albero",
    "anchor": "Ancora",
    "mulio": "Mulino",
}


def clean(value):
    """Trim strings, drop placeholders, turn whole floats into ints."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    value = re.sub(r"\s+", " ", str(value)).strip()
    return None if value.lower() in PLACEHOLDERS else value


def normalize(value, mapping):
    value = clean(value)
    if value is None:
        return None, None
    fixed = mapping.get(str(value).lower(), value)
    if fixed == value and isinstance(value, str):
        fixed = value[0].upper() + value[1:]  # "wet" -> "Wet"
    return fixed, (value if fixed != value else None)


def lat_lng(a, b):
    """Return (lat, lng) for the Venice lagoon whichever order the sheet used."""
    try:
        a, b = float(a), float(b)
    except (TypeError, ValueError):
        return None, None
    if 44 < a < 47 and 11 < b < 14:
        return a, b
    if 44 < b < 47 and 11 < a < 14:
        return b, a
    return None, None


def sito_key(catalog):
    """'0127-1-1' -> '127.1.1', to match QGIS 'Sito' ids."""
    if catalog is None:
        return None
    parts = str(catalog).split("-")
    parts[0] = parts[0].lstrip("0") or "0"
    return ".".join(parts)


def compact(d):
    return {k: v for k, v in d.items() if v is not None}


def load_gis():
    wb = openpyxl.load_workbook(GIS_XLSX, data_only=True)
    features, seen = {}, set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] != "Feature":  # skips blank and repeated header rows
                continue
            lng, lat = float(row[15]), float(row[16])
            props = tuple(clean(v) for v in row[1:14])
            if (props, round(lat, 7), round(lng, 7)) in seen:
                continue
            seen.add((props, round(lat, 7), round(lng, 7)))
            simboli, simboli_orig = normalize(row[2], TYPES)
            key = "gis-%04d" % (len(features) + 1)
            features[key] = compact({
                "sito": clean(row[1]),
                "simboli": simboli,
                "simboli_original": simboli_orig,
                "epoche": clean(row[3]),
                "name": clean(row[4]),
                "struttura": clean(row[5]),
                "rilievo": clean(row[6]),
                "rilevamento": clean(row[7]),
                "coordinate": clean(row[8]),
                "profondita": clean(row[9]),
                "datazione": clean(row[10]),
                "c14": clean(row[11]),
                "numerazione": clean(row[12]),
                "orientamento": clean(row[13]),
                "layer": ws.title,
                "lat": lat,
                "lng": lng,
            })
    return features


def build():
    for path in (SITE_XLSX, GIS_XLSX):
        if not os.path.exists(path):
            sys.exit("Missing input: %s" % path)

    gis = load_gis()
    gis_by_sito = {}
    for key, f in gis.items():
        if "sito" in f:
            gis_by_sito.setdefault(str(f["sito"]), []).append(key)

    wb = openpyxl.load_workbook(SITE_XLSX, data_only=True)
    data, groups = {}, {"Site": {"members": {}}, "Layers": {"members": {}}}
    warnings = Counter()

    sites = [r for r in wb["SiteForm"].iter_rows(min_row=2, values_only=True)
             if any(clean(c) is not None for c in r)]
    for i, r in enumerate(sites, start=1):
        key = "site-%04d" % i
        catalog = clean(r[4])
        catalog = str(catalog) if catalog is not None else None
        lat, lng = lat_lng(r[7], r[8])
        if lat is None:
            warnings["site without valid coordinates"] += 1
        epoch, epoch_orig = normalize(r[18], EPOCHS)
        stype, type_orig = normalize(r[22], TYPES)
        sito = sito_key(catalog)
        gis_keys = gis_by_sito.get(sito, [])
        place = next((gis[k]["name"] for k in gis_keys if "name" in gis[k]), None)
        site = compact({
            "catalog": catalog,
            "sito": sito,
            "name": place,
            "year": clean(r[1]),
            "archaeologist": clean(r[2]),
            "scimanage": clean(r[3]),
            "sopr": clean(r[5]),
            "center": normalize(r[6], {"north lagoon": "North Lagoon",
                                       "south lagoon": "South Lagoon"})[0],
            "lat": lat,
            "long": lng,
            "type1": clean(r[9]),
            "type2": normalize(r[10], {})[0],
            "objfind": clean(r[11]),
            "analysis": clean(r[12]),
            "drawmap": clean(r[13]),
            "drawsec": clean(r[14]),
            "photo": clean(r[15]),
            "digipho": clean(r[16]),
            "finds": clean(r[17]),
            "epoch": epoch,
            "epoch_original": epoch_orig,
            "from": clean(r[19]),
            "to": clean(r[20]),
            "interp": clean(r[21]),
            "type": stype,
            "type_original": type_orig,
            "gis": {k: True for k in gis_keys} or None,
        })
        data[key] = {
            "birth_certificate": compact({
                "birthID": key,
                "ckID": key,
                "dor": clean(r[0]) or IMPORT_DATE,
                "recorder": RECORDER,
                "type": "Site",
                "group": "Site",
                "lat": lat,
                "long": lng,
            }),
            "data": dict(site, ckID=key),
        }
        groups["Site"]["members"][key] = key
        for k in gis_keys:
            gis[k]["site_id"] = key

    by_catalog = {}
    for key, rec in data.items():
        if "sito" in rec["data"]:
            by_catalog.setdefault(rec["data"]["sito"], key)

    for r in wb["USLayers"].iter_rows(min_row=2, values_only=True):
        if clean(r[2]) is None:
            continue
        site_no, us = clean(r[1]), clean(r[2])
        key = "layer-%s-%s" % (site_no, us)
        site_id = by_catalog.get(sito_key(site_no))
        if site_id is None:
            warnings["layer whose site is not in SiteForm"] += 1
        data[key] = {
            "birth_certificate": {
                "birthID": key, "ckID": key, "dor": clean(r[0]) or IMPORT_DATE,
                "recorder": RECORDER, "type": "Layers",
            },
            "data": compact({
                "ckID": key,
                "site_id": site_id,
                "sitenumber": site_no,
                "us": us,
                "usnumber": us,
                "description": clean(r[3]),
                "coveredby": clean(r[4]),
                "covering": clean(r[5]),
            }),
        }
        groups["Layers"]["members"][key] = key
        if site_id:
            data[site_id]["data"].setdefault("USlayers", {})[key] = {
                "site_ckId": key, "name": "US %s" % us}

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump({"data": data, "groups": groups, "gis": {"features": gis}},
                  fh, ensure_ascii=False, indent=1, sort_keys=True)

    linked = sum(1 for f in gis.values() if "site_id" in f)
    print("sites:        %d" % len(groups["Site"]["members"]))
    print("layers:       %d" % len(groups["Layers"]["members"]))
    print("gis features: %d (%d linked to a site)" % (len(gis), linked))
    for msg, n in warnings.items():
        print("warning: %d %s" % (n, msg))
    print("wrote %s (%.0f KB)" % (os.path.relpath(OUT, ROOT), os.path.getsize(OUT) / 1024))


if __name__ == "__main__":
    build()
